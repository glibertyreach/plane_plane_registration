"""Check both registration decks against their content JSONs: every non-notes text string in the JSON
appears on its slide in the .pptx (via markitdown), and every slide's speaker notes equal the JSON notes.
Then check the cost figures on the build deck against docs/procedures/build/costs.py, their single source.

Run from the repository root:  python3 docs/presentations/archive/check_registration_content.py"""
import json, re, subprocess, sys
from pathlib import Path
from pptx import Presentation

DECKS = [  # (content JSON, built deck)
    ("docs/presentations/build/registration_build_deck_content.json", "docs/presentations/registration_procurement_build.pptx"),
    ("docs/presentations/build/registration_procedure_deck_content.json", "docs/presentations/registration_test_procedure.pptx"),
]
SKIP_KEYS = {"id", "layout", "section", "notes", "icon", "image", "image2", "n"}  # not slide text

def norm(t):
    t = t.replace(" ", " ").replace("\\", "")      # non-breaking space, markdown escapes
    t = re.sub(r"\s+", " ", t)
    return t.strip()

def strings(obj):
    if isinstance(obj, str):
        yield obj
    elif isinstance(obj, list):
        for v in obj: yield from strings(v)
    elif isinstance(obj, dict):
        for k, v in obj.items():
            if k not in SKIP_KEYS: yield from strings(v)

def check(content_json, pptx):
    """Return the number of problems found in one deck."""
    md = subprocess.run(["markitdown", pptx], capture_output=True, text=True, check=True).stdout
    parts = re.split(r"<!-- Slide number: (\d+) -->", md)
    # markitdown appends "### Notes:" to each slide block; keep the slide body and the notes apart
    slide_text, notes_text = {}, {}
    for i in range(1, len(parts), 2):
        body, _, notes = parts[i + 1].partition("### Notes:")
        slide_text[int(parts[i])], notes_text[int(parts[i])] = norm(body), norm(notes)
    slides = json.load(open(content_json))["slides"]
    prs = Presentation(pptx)
    bad = 0
    assert len(slides) == len(prs.slides) == len(slide_text), (len(slides), len(prs.slides), len(slide_text))
    for i, s in enumerate(slides, 1):
        text = slide_text[i]
        for t in strings({k: v for k, v in s.items()}):
            if norm(t) not in text:
                bad += 1; print(f"slide {i} ({s['id']}): MISSING {t!r}")
        notes = prs.slides[i - 1].notes_slide.notes_text_frame.text if prs.slides[i - 1].has_notes_slide else ""
        if not notes.strip() or norm(notes) != norm(s["notes"]):
            bad += 1; print(f"slide {i} ({s['id']}): NOTES differ or empty")
        # notes text must not be on the slide itself
        if norm(s["notes"]) != notes_text[i]:
            bad += 1; print(f"slide {i}: markitdown notes differ from JSON notes")
        if norm(s["notes"])[:60] in text:
            bad += 1; print(f"slide {i}: notes text found on the slide body")
    print(f"{pptx}: checked {len(slides)} slides; problems: {bad}")
    return bad


# ---- cost consistency: the build deck's figures against costs.py ---------------------------------------
sys.path.insert(0, str(Path("docs/procedures/build").resolve()))
import costs  # noqa: E402  (the single source of the cost figures)

COST_HEADER = "Estimated cost (USD)"
FROM_STAGE_1_HEADER = "From stage 1"


def find_item(row_name, items):
    """The costs.py item whose name the row name starts with (or the reverse); failing that, the only item whose
    name contains the row name's first word (the deck says 'Straightedge 300 mm and feeler gauges', costs.py
    'Precision straightedge, 300 mm, and feeler gauges'). Returns (item, how) or (None, None)."""
    low = row_name.lower()
    for it in items:
        n = it.name.lower()
        if low.startswith(n) or n.startswith(low):
            return it, "name prefix"
    word = re.split(r"\W+", low)[0]
    hits = [it for it in items if word in re.split(r"\W+", it.name.lower())]
    return (hits[0], f"first word '{word}'") if len(hits) == 1 else (None, None)


def check_costs(content_json):
    """Compare every cost string of the build deck's cost and buy_list slides with costs.py; return the problems found."""
    slides = {s["id"]: s for s in json.load(open(content_json))["slides"]}
    bad = 0
    for slide_id, items in (("cost", costs.BUILD_ITEMS), ("buy_list", costs.BUILD_ITEMS + costs.BUY_ITEMS)):
        table = slides[slide_id]["table"]
        ci = table["header"].index(COST_HEADER)
        si = table["header"].index(FROM_STAGE_1_HEADER) if FROM_STAGE_1_HEADER in table["header"] else None
        for row in table["rows"]:
            item, how = find_item(row[0], items)
            if item is None:
                print(f"  {slide_id}: {row[0]!r}: {row[ci]!r}: no costs.py item, not compared")
                continue
            want = costs.usd_range(item.low, item.high)
            ok = row[ci] == want
            if si is not None:
                stage1 = "yes" if item.in_hand_from_stage_1 else "no"
                ok = ok and row[si] == stage1
            bad += 0 if ok else 1
            print(f"  {slide_id}: {row[0]!r} ~ {item.name!r} ({how}): deck {row[ci]!r}"
                  f"{'' if si is None else ', from stage 1 ' + repr(row[si])} vs costs.py {want!r}: {'ok' if ok else 'MISMATCH'}")
    j = costs.as_json()
    stats = slides["cost"]["stats"]
    for stat, key in zip(stats, ("new_spending_if_stage_1_done", "total")):
        ok = stat["value"] == j[key]
        bad += 0 if ok else 1
        print(f"  cost stat {stat['value']!r} vs costs.as_json()[{key!r}] = {j[key]!r}: {'ok' if ok else 'MISMATCH'}")
    print(f"cost check: problems: {bad}")
    return bad


total = sum(check(content_json, pptx) for content_json, pptx in DECKS)
print(f"total problems: {total}")
print("cost consistency (build deck vs docs/procedures/build/costs.py):")
cost_bad = check_costs(DECKS[0][0])
sys.exit(1 if total or cost_bad else 0)
